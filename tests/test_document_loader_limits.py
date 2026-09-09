from __future__ import annotations

import io
import zipfile
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest

import services.document_loader as document_loader
from services.document_loader import (
    DocumentDownloadError,
    DocumentLoader,
    DocumentParseError,
)


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)
    return output.getvalue()


def _loader_with_response(response: dict[str, object]) -> DocumentLoader:
    loader = DocumentLoader.__new__(DocumentLoader)
    loader.bucket = "uploads"
    loader.client = MagicMock()
    loader.client.get_object.return_value = response
    return loader


def test_s3_download_is_bounded_and_body_is_closed(monkeypatch):
    monkeypatch.setattr(document_loader, "MAX_FILE_SIZE_BYTES", 4)
    body = io.BytesIO(b"12345")
    loader = _loader_with_response({"Body": body, "ContentType": "text/plain"})

    with pytest.raises(DocumentDownloadError, match="4-byte download limit"):
        loader.fetch_text("oversized.txt")

    assert body.closed is True


def test_declared_s3_size_is_rejected_without_reading_body(monkeypatch):
    monkeypatch.setattr(document_loader, "MAX_FILE_SIZE_BYTES", 4)
    body = MagicMock()
    loader = _loader_with_response(
        {"Body": body, "ContentLength": 5, "ContentType": "text/plain"}
    )

    with pytest.raises(DocumentDownloadError, match="4-byte download limit"):
        loader.fetch_text("oversized.txt")

    body.read.assert_not_called()
    body.close.assert_called_once_with()


def test_private_parse_boundary_rejects_oversized_bytes(monkeypatch):
    monkeypatch.setattr(document_loader, "MAX_FILE_SIZE_BYTES", 4)
    loader = DocumentLoader.__new__(DocumentLoader)

    with pytest.raises(DocumentParseError, match="4-byte parse limit"):
        loader._extract_text(b"12345", "oversized.txt", "text/plain")


def test_plain_text_is_capped_without_corrupting_chunk_boundaries(monkeypatch):
    monkeypatch.setattr(document_loader, "MAX_EXTRACTED_TEXT_CHARS", 65_536)
    loader = DocumentLoader.__new__(DocumentLoader)
    # The first byte of `ü` lands at the end of the decoder's 64 KiB block.
    payload = b"a" * 65_535 + "üç".encode("utf-8")

    result = loader._extract_text(payload, "notes.txt", "text/plain")

    assert result == "a" * 65_535 + "ü"
    assert len(result) == 65_536


def test_office_archive_entry_size_is_checked_before_parser(monkeypatch):
    monkeypatch.setattr(document_loader, "OFFICE_MAX_ENTRY_SIZE_BYTES", 10)
    payload = _zip_bytes({"word/document.xml": b"x" * 11})

    with pytest.raises(DocumentParseError, match="archive entry size limit exceeded"):
        document_loader._validate_office_archive(payload, "DOCX")


def test_office_archive_total_uncompressed_size_is_bounded(monkeypatch):
    monkeypatch.setattr(document_loader, "OFFICE_MAX_ENTRY_SIZE_BYTES", 20)
    monkeypatch.setattr(document_loader, "OFFICE_MAX_TOTAL_UNCOMPRESSED_BYTES", 15)
    payload = _zip_bytes({"one.xml": b"a" * 8, "two.xml": b"b" * 8})

    with pytest.raises(DocumentParseError, match="uncompressed size limit exceeded"):
        document_loader._validate_office_archive(payload, "PPTX")


def test_pdf_page_limit_stops_before_text_extraction(monkeypatch):
    first_page = SimpleNamespace(extract_text=MagicMock(return_value="first"))
    second_page = SimpleNamespace(extract_text=MagicMock(return_value="second"))
    monkeypatch.setattr(document_loader, "PDF_MAX_PAGES", 1)
    monkeypatch.setattr(
        document_loader,
        "PdfReader",
        lambda _stream: SimpleNamespace(pages=[first_page, second_page]),
    )

    with pytest.raises(DocumentParseError, match=r"PDF page limit exceeded \(2 > 1\)"):
        DocumentLoader._read_pdf(b"fake-pdf")

    first_page.extract_text.assert_not_called()
    second_page.extract_text.assert_not_called()


def test_docx_paragraph_limit_is_checked_before_collecting_text(monkeypatch):
    paragraphs = [SimpleNamespace(text="one"), SimpleNamespace(text="two")]
    monkeypatch.setattr(document_loader, "DOCX_MAX_PARAGRAPHS", 1)
    monkeypatch.setattr(document_loader, "_validate_office_archive", lambda *_args: ())
    monkeypatch.setattr(
        document_loader, "Document", lambda _stream: SimpleNamespace(paragraphs=paragraphs)
    )

    with pytest.raises(DocumentParseError, match="DOCX paragraph limit exceeded"):
        DocumentLoader._read_docx(b"fake-docx")


def test_xlsx_dimension_limit_closes_workbook_and_stream(monkeypatch):
    sheet = SimpleNamespace(title="TooManyRows", max_row=2, max_column=1)
    workbook = SimpleNamespace(worksheets=[sheet], close=MagicMock())
    monkeypatch.setattr(document_loader, "XLSX_MAX_ROWS_PER_SHEET", 1)
    monkeypatch.setattr(document_loader, "_validate_office_archive", lambda *_args: ())

    import openpyxl

    monkeypatch.setattr(openpyxl, "load_workbook", lambda **_kwargs: workbook)

    with pytest.raises(DocumentParseError, match="XLSX row limit exceeded"):
        DocumentLoader._read_xlsx(b"fake-xlsx")

    workbook.close.assert_called_once_with()


def test_xlsx_reads_only_validated_dimensions_and_closes_workbook(monkeypatch):
    sheet = SimpleNamespace(title="Sheet 1", max_row=2, max_column=2)
    sheet.iter_rows = MagicMock(return_value=iter([("A", "B"), (1, 2)]))
    workbook = SimpleNamespace(worksheets=[sheet], close=MagicMock())
    monkeypatch.setattr(document_loader, "_validate_office_archive", lambda *_args: ())

    import openpyxl

    monkeypatch.setattr(openpyxl, "load_workbook", lambda **_kwargs: workbook)

    result = DocumentLoader._read_xlsx(b"fake-xlsx")

    assert result == "[sheet] Sheet 1\nA | B\n1 | 2"
    sheet.iter_rows.assert_called_once_with(
        min_row=1, max_row=2, min_col=1, max_col=2, values_only=True
    )
    workbook.close.assert_called_once_with()


def test_pptx_slide_limit_is_checked_from_zip_before_parser(monkeypatch):
    monkeypatch.setattr(document_loader, "PPTX_MAX_SLIDES", 1)
    payload = _zip_bytes(
        {
            "ppt/slides/slide1.xml": b"<slide />",
            "ppt/slides/slide2.xml": b"<slide />",
        }
    )

    with pytest.raises(DocumentParseError, match=r"PPTX slide limit exceeded \(2 > 1\)"):
        DocumentLoader._read_pptx(payload)


def test_pptx_shape_limit_is_enforced(monkeypatch):
    shapes = [SimpleNamespace(text="one"), SimpleNamespace(text="two")]
    slides = [[*shapes]]
    presentation = SimpleNamespace(
        slides=[SimpleNamespace(shapes=slide_shapes) for slide_shapes in slides]
    )
    monkeypatch.setattr(document_loader, "PPTX_MAX_SHAPES", 1)
    monkeypatch.setattr(document_loader, "_validate_office_archive", lambda *_args: ())

    import pptx

    monkeypatch.setattr(pptx, "Presentation", lambda _stream: presentation)

    with pytest.raises(DocumentParseError, match="PPTX shape limit exceeded"):
        DocumentLoader._read_pptx(b"fake-pptx")
